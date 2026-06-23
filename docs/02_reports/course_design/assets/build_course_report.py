from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

import build_final_report as base


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "docs" / "output"
DOCX_OUT = OUTPUT / "MiniOS项目报告_课程设计版_张晓文_20231071455.docx"


def body_paragraph(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.32
    p.paragraph_format.space_after = Pt(3)
    base.add_inline_markup(p, text, size=10.5)
    return p


def list_item(doc, text, numbered=False):
    p = doc.add_paragraph(style="List Number" if numbered else "List Bullet")
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(2)
    base.add_inline_markup(p, text, size=10.5)
    return p


def configure_course_styles(doc):
    base.configure_styles(doc)
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)

    settings = {
        "Heading 1": (15, 15, 8),
        "Heading 2": (12.5, 11, 5),
        "Heading 3": (11, 8, 4),
    }
    for name, (size, before, after) in settings.items():
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_front_matter(doc):
    cover = doc.sections[0]
    base.set_section_geometry(cover)
    cover.top_margin = Cm(2.8)
    vertical = OxmlElement("w:vAlign")
    vertical.set(qn("w:val"), "center")
    cover._sectPr.append(vertical)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(38)
    base.set_run_font(
        p.add_run("操作系统课程设计报告"),
        name="Microsoft YaHei", size=18, bold=True,
    )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(52)
    base.set_run_font(
        p.add_run(base.TITLE),
        name="Microsoft YaHei", size=21, bold=True,
    )

    for label, value in (
        ("学生姓名", base.AUTHOR),
        ("学号", base.STUDENT_ID),
        ("日期", "2026 年 6 月"),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
        base.set_run_font(p.add_run(f"{label}：{value}"), name="Microsoft YaHei", size=12)

    front = doc.add_section(WD_SECTION.NEW_PAGE)
    base.set_section_geometry(front)
    front.top_margin = Cm(2.2)
    front.bottom_margin = Cm(2.2)
    front.header.is_linked_to_previous = False
    front.footer.is_linked_to_previous = False

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    base.set_run_font(p.add_run("摘要"), name="Microsoft YaHei", size=15, bold=True)
    abstract = (
        "本项目在自研 RISC-V RV64I 模拟平台 cemu 上设计并实现教学型操作系统 MiniOS。"
        "系统支持两阶段启动、M/S/U 特权级、Sv39 虚拟内存、系统调用、进程管理、"
        "RR/FCFS 调度、持久化 MiniFS v2、ELF64 用户程序和交互式 Shell。"
        "项目将硬件单元测试、镜像工具测试和 7 个系统级用例统一纳入 CTest，"
        "最近一次回归测试 99 项全部通过，形成了从模拟硬件、Bootloader、内核到"
        "用户程序的完整运行闭环。"
    )
    body_paragraph(doc, abstract)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    base.set_run_font(p.add_run("关键词："), name="Microsoft YaHei", size=10.5, bold=True)
    base.set_run_font(
        p.add_run("RISC-V；MiniOS；Sv39；进程调度；MiniFS；ELF"),
        name="Microsoft YaHei", size=10.5,
    )

    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    base.set_run_font(p.add_run("目录"), name="Microsoft YaHei", size=15, bold=True)
    base.add_toc(doc.add_paragraph())

    body = doc.add_section(WD_SECTION.NEW_PAGE)
    base.set_section_geometry(body)
    body.top_margin = Cm(2.2)
    body.bottom_margin = Cm(2.0)
    body.left_margin = Cm(2.4)
    body.right_margin = Cm(2.4)
    body.header.is_linked_to_previous = False
    body.footer.is_linked_to_previous = False
    base.set_page_numbering(body, "decimal", 1)
    header = body.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    base.set_run_font(
        header.add_run("MiniOS 课程设计报告"),
        name="Microsoft YaHei", size=8.5, color=base.GRAY,
    )
    base.add_page_field(body.footer.paragraphs[0])


def build():
    OUTPUT.mkdir(parents=True, exist_ok=True)

    # 保留完整内容，只调整视觉密度和分页策略。
    base.DOCX_OUT = DOCX_OUT
    base.CHAPTER_PAGE_BREAK = False
    base.REFERENCE_PAGE_BREAK = False
    base.OWN_PAGE_FIGURES = set()
    base.FONT_CN = "Microsoft YaHei"
    base.FONT_HEADING = "Microsoft YaHei"
    base.add_body_paragraph = body_paragraph
    base.add_list_item = list_item

    diagrams = base.build_diagrams()
    terminal = base.capture_terminal()

    doc = Document()
    configure_course_styles(doc)
    add_front_matter(doc)
    base.parse_body(doc, diagrams, terminal)
    base.set_update_fields(doc)
    doc.core_properties.title = base.TITLE
    doc.core_properties.author = base.AUTHOR
    doc.core_properties.subject = "操作系统课程设计报告（简洁排版版）"
    doc.save(DOCX_OUT)
    print(DOCX_OUT)


if __name__ == "__main__":
    build()
