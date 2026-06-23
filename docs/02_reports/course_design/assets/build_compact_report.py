from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

import build_final_report as base
import build_course_report as course


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "docs" / "output"
DOCX_OUT = OUTPUT / "MiniOS项目报告_精简排版版_张晓文_20231071455.docx"


def add_front_matter(doc):
    cover = doc.sections[0]
    base.set_section_geometry(cover)
    cover.top_margin = Cm(2.8)
    vertical = OxmlElement("w:vAlign")
    vertical.set(qn("w:val"), "center")
    cover._sectPr.append(vertical)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(38)
    base.set_run_font(
        p.add_run("操作系统课程设计报告"),
        name="Microsoft YaHei", size=18, bold=True,
    )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(52)
    base.set_run_font(
        p.add_run(base.TITLE),
        name="Microsoft YaHei", size=21, bold=True,
    )

    for label, value in (
        ("学生姓名", base.AUTHOR),
        ("学号", base.STUDENT_ID),
        ("日期", "2026 年 6 月"),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
        base.set_run_font(
            p.add_run(f"{label}：{value}"),
            name="Microsoft YaHei", size=12,
        )

    toc_section = doc.add_section(WD_SECTION.NEW_PAGE)
    base.set_section_geometry(toc_section)
    toc_section.top_margin = Cm(2.2)
    toc_section.bottom_margin = Cm(2.2)
    toc_section.header.is_linked_to_previous = False
    toc_section.footer.is_linked_to_previous = False

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    base.set_run_font(
        p.add_run("目录"),
        name="Microsoft YaHei", size=15, bold=True,
    )
    base.add_toc(doc.add_paragraph())

    body = doc.add_section(WD_SECTION.NEW_PAGE)
    base.set_section_geometry(body)
    body.top_margin = Cm(2.2)
    body.bottom_margin = Cm(2.0)
    body.left_margin = Cm(2.4)
    body.right_margin = Cm(2.4)
    body.header.is_linked_to_previous = False
    body.footer.is_linked_to_previous = False
    base.set_page_numbering(body, "decimal", 1)
    header = body.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    base.set_run_font(
        header.add_run("MiniOS 课程设计报告"),
        name="Microsoft YaHei", size=8.5, color=base.GRAY,
    )
    base.add_page_field(body.footer.paragraphs[0])


def build():
    OUTPUT.mkdir(parents=True, exist_ok=True)

    base.DOCX_OUT = DOCX_OUT
    base.CHAPTER_PAGE_BREAK = False
    base.REFERENCE_PAGE_BREAK = False
    base.OWN_PAGE_FIGURES = set()
    base.INCLUDE_REFERENCES = False
    base.TOC_LEVELS = "1-1"
    base.FIGURE_SCALE = 0.85
    base.TABLE_FONT_SIZE = 8.5
    base.TABLE_CELL_MARGIN_TOP = 55
    base.TABLE_CELL_MARGIN_BOTTOM = 55
    base.FONT_CN = "Microsoft YaHei"
    base.FONT_HEADING = "Microsoft YaHei"
    base.add_body_paragraph = course.body_paragraph
    base.add_list_item = course.list_item

    diagrams = base.build_diagrams()
    terminal = base.capture_terminal()

    doc = Document()
    course.configure_course_styles(doc)
    add_front_matter(doc)
    base.parse_body(doc, diagrams, terminal)
    base.set_update_fields(doc)
    doc.core_properties.title = base.TITLE
    doc.core_properties.author = base.AUTHOR
    doc.core_properties.subject = "操作系统课程设计报告（精简排版版）"
    doc.save(DOCX_OUT)
    print(DOCX_OUT)


if __name__ == "__main__":
    build()
