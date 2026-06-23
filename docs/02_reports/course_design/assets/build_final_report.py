from __future__ import annotations

import html
import re
import subprocess
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "docs" / "项目报告.md"
ASSETS = ROOT / "docs" / "report_assets" / "final"
OUTPUT = ROOT / "docs" / "output"
DOCX_OUT = OUTPUT / "MiniOS项目报告_张晓文_20231071455.docx"

TITLE = "基于自研 RV64I 模拟平台的 MiniOS 操作系统设计与实现"
SHORT_TITLE = "MiniOS 操作系统设计与实现"
AUTHOR = "张晓文"
STUDENT_ID = "20231071455"
DATE_TEXT = "2026 年 6 月"
CHAPTER_PAGE_BREAK = True
REFERENCE_PAGE_BREAK = True
OWN_PAGE_FIGURES = {"【图 4 占位", "【图 5 占位"}
INCLUDE_REFERENCES = True
TOC_LEVELS = "1-3"
FIGURE_SCALE = 1.0
TABLE_FONT_SIZE = 8.7
TABLE_CELL_MARGIN_TOP = 90
TABLE_CELL_MARGIN_BOTTOM = 90

FONT_CN = "SimSun"
FONT_HEADING = "SimHei"
FONT_LATIN = "Times New Roman"
FONT_CODE = "Consolas"
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(90, 90, 90)


def set_run_font(run, name=FONT_CN, size=12, bold=False, color=BLACK, italic=False):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    fonts.set(qn("w:ascii"), FONT_LATIN if name == FONT_CN else name)
    fonts.set(qn("w:hAnsi"), FONT_LATIN if name == FONT_CN else name)
    fonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table, top=12, bottom=12, inside=4):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge, size in (("top", top), ("bottom", bottom), ("insideH", inside)):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:color"), "000000")
    for edge in ("left", "right", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])
    set_run_font(run, size=9)


def set_page_numbering(section, fmt="decimal", start=1):
    sect_pr = section._sectPr
    node = sect_pr.find(qn("w:pgNumType"))
    if node is None:
        node = OxmlElement("w:pgNumType")
        sect_pr.append(node)
    node.set(qn("w:fmt"), fmt)
    node.set(qn("w:start"), str(start))


def set_section_geometry(section):
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.4)
    section.footer_distance = Cm(1.4)


def add_toc(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f'TOC \\o "{TOC_LEVELS}" \\h \\z \\u'
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "目录将在导出 PDF 前自动更新。"
    sep_run = OxmlElement("w:r")
    sep_run.append(text)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, sep_run, end])


def add_inline_markup(paragraph, text, size=12, color=BLACK):
    pattern = re.compile(r"(\*\*.*?\*\*|`.*?`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            set_run_font(paragraph.add_run(text[pos:match.start()]), size=size, color=color)
        token = match.group(0)
        if token.startswith("**"):
            set_run_font(paragraph.add_run(token[2:-2]), size=size, bold=True, color=color)
        else:
            set_run_font(paragraph.add_run(token[1:-1]), name=FONT_CODE, size=size - 0.5, color=color)
        pos = match.end()
    if pos < len(text):
        set_run_font(paragraph.add_run(text[pos:]), size=size, color=color)


def add_body_paragraph(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt = p.paragraph_format
    fmt.first_line_indent = Cm(0.74)
    fmt.line_spacing = 1.5
    fmt.space_after = Pt(0)
    add_inline_markup(p, text)
    return p


def add_list_item(doc, text, numbered=False):
    p = doc.add_paragraph(style="List Number" if numbered else "List Bullet")
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(0)
    add_inline_markup(p, text)
    return p


def create_numbering(doc, start_value=1):
    root = doc.part.numbering_part.element
    abstract_ids = [
        int(x.get(qn("w:abstractNumId")))
        for x in root.findall(qn("w:abstractNum"))
    ]
    num_ids = [int(x.get(qn("w:numId"))) for x in root.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), str(start_value))
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), "decimal")
    text = OxmlElement("w:lvlText")
    text.set(qn("w:val"), "%1.")
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "space")
    ppr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    ppr.append(ind)
    level.extend([start, fmt, text, suff, ppr])
    abstract.append(level)
    first_num = root.find(qn("w:num"))
    if first_num is None:
        root.append(abstract)
    else:
        root.insert(list(root).index(first_num), abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abstract_id))
    num.append(ref)
    root.append(num)
    return num_id


def add_numbered_group(doc, items, start_value=1):
    num_id = create_numbering(doc, start_value)
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(0)
        ppr = p._p.get_or_add_pPr()
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num = OxmlElement("w:numId")
        num.set(qn("w:val"), str(num_id))
        num_pr.extend([ilvl, num])
        ppr.append(num_pr)
        add_inline_markup(p, item)


def add_code_block(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Cm(15.6)
    set_cell_fill(cell, "F2F2F2")
    set_cell_margins(cell, top=100, start=140, bottom=100, end=140)
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.05
    p.paragraph_format.space_after = Pt(0)
    for idx, line in enumerate(text.rstrip().splitlines()):
        if idx:
            p.add_run().add_break()
        set_run_font(p.add_run(line), name=FONT_CODE, size=8.3)
    table.style = "Table Grid"
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def add_markdown_table(doc, rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=1, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_repeat_header(table.rows[0])
    for c in range(cols):
        cell = table.rows[0].cells[c]
        text = rows[0][c] if c < len(rows[0]) else ""
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(
            cell, top=TABLE_CELL_MARGIN_TOP, bottom=TABLE_CELL_MARGIN_BOTTOM
        )
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_inline = p.add_run(re.sub(r"\*\*|`", "", text))
        set_run_font(set_inline, name=FONT_HEADING, size=9, bold=True)
    for row in rows[1:]:
        cells = table.add_row().cells
        for c in range(cols):
            text = row[c] if c < len(row) else ""
            cells[c].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(
                cells[c], top=TABLE_CELL_MARGIN_TOP,
                bottom=TABLE_CELL_MARGIN_BOTTOM
            )
            p = cells[c].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.15
            add_inline_markup(p, text, size=TABLE_FONT_SIZE)
    set_table_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def dot_escape(text):
    return html.escape(text, quote=True)


def render_dot(name, source):
    ASSETS.mkdir(parents=True, exist_ok=True)
    dot_path = ASSETS / f"{name}.dot"
    png_path = ASSETS / f"{name}.png"
    dot_path.write_text(source, encoding="utf-8")
    subprocess.run(
        ["dot", "-Tpng", "-Gdpi=180", str(dot_path), "-o", str(png_path)],
        check=True,
    )
    return png_path


def diagram_source(title, body, rankdir="TB"):
    return f'''digraph G {{
graph [bgcolor="white", pad=0.18, nodesep=0.35, ranksep=0.52, rankdir={rankdir}];
node [shape=box, style="rounded,filled", fillcolor="#F2F2F2", color="black",
      fontname="Noto Sans SC", fontsize=12, margin="0.18,0.10"];
edge [color="black", arrowsize=0.75, penwidth=1.1, fontname="Noto Sans SC", fontsize=10];
{body}
}}'''


def build_diagrams():
    diagrams = {}
    diagrams[1] = render_dot("fig1_overall", diagram_source(
        "MiniOS 总体分层架构",
        '''
host [label="宿主环境\\nWSL / CMake / Make / RISC-V 工具链"];
cemu [label="cemu 模拟硬件层\\nCPU / CSR / MMU / Bus / DRAM / 外设"];
boot [label="两阶段启动\\nboot.bin → kernel.bin → S-mode"];
kernel [label="MiniOS 内核层\\nTrap / Syscall / VM / Scheduler / MiniFS"];
user [label="用户程序层\\nELF / libc / Shell / 命令与测试程序"];
host -> cemu -> boot -> kernel -> user;
'''))
    diagrams[2] = render_dot("fig2_cemu", diagram_source(
        "cemu 模拟器模块关系",
        '''
main [label="main.cpp\\n加载 boot.bin 与 disk.img"];
cpu [label="CPU Core\\nPC / GPR / Privilege"];
ins [label="Instruction Executor\\nRV64I 译码与执行"];
csr [label="CSR\\n异常 / 中断 / 委托"];
mmu [label="MMU\\nSv39 地址翻译"];
bus [label="Bus\\n物理地址与 MMIO 路由"];
dram [label="DRAM\\n128 MiB"];
uart [label="UART"];
clint [label="CLINT"];
plic [label="PLIC"];
block [label="Block Device\\n8 MiB disk.img"];
main -> cpu; cpu -> ins; cpu -> csr; cpu -> mmu; mmu -> bus; ins -> bus;
bus -> dram; bus -> uart; bus -> clint; bus -> plic; bus -> block;
{rank=same; ins; csr; mmu}
{rank=same; dram; uart; clint; plic; block}
'''))
    diagrams[3] = render_dot("fig3_instruction", diagram_source(
        "CPU 单条指令执行流程",
        '''
start [label="检查停机状态\\n推进 CLINT"];
irq [label="检查待处理中断"];
fetch [label="按 PC 取指"];
translate [label="MMU 地址翻译\\n权限检查"];
decode [label="opcode / funct3 / funct7 译码"];
execute [label="ALU / 访存 / 跳转 / CSR"];
update [label="更新 PC"];
trap [label="异常或中断？", shape=diamond, fillcolor="white"];
state [label="写 EPC / CAUSE / TVAL\\n切换特权级并跳转 TVEC"];
start -> irq -> fetch -> translate -> decode -> execute -> update -> trap;
trap -> start [label="否"];
trap -> state [label="是"];
state -> start;
'''))
    diagrams[4] = render_dot("fig4_boot", diagram_source(
        "MiniOS 两阶段启动流程",
        '''
cemu [label="cemu 仅加载 boot.bin\\n0x80000000，M-mode"];
header [label="Bootloader 读取扇区 15360\\n校验 magic / version / size / entry"];
load [label="读取 kernel.bin\\n扇区 15361-16383"];
checksum [label="计算并核对 32 位 checksum"];
copy [label="搬运内核到 0x80200000\\n跳转 _start"];
init [label="设置栈 / 清零 BSS\\n配置 mtvec、stvec 与委托"];
mret [label="mstatus.MPP=S\\nmepc=kernel_main\\nmret"];
shell [label="初始化内核子系统\\n进入 U-mode Shell"];
cemu -> header -> load -> checksum -> copy -> init -> mret -> shell;
'''))
    diagrams[5] = render_dot("fig5_trap", diagram_source(
        "Trap、系统调用与调度链路",
        '''
event [label="U-mode ecall\\n或 CLINT 定时器中断"];
entry [label="trap_entry\\n切换内核栈并保存 32 个寄存器"];
handler [label="trap_handler\\n读取 scause / sepc / stval"];
kind [label="事件类型", shape=diamond, fillcolor="white"];
sys [label="syscall_dispatch\\na7=编号，a0-a2=参数"];
timer [label="timer_handle\\n重设 mtimecmp"];
sched [label="sched_tick\\nRR 时间片检查"];
switch [label="保存/恢复 trap context\\n切换 SATP 与任务"];
returnn [label="恢复寄存器并 sret"];
event -> entry -> handler -> kind;
kind -> sys [label="ecall"];
kind -> timer [label="timer"];
timer -> sched -> switch -> returnn;
sys -> returnn;
'''))
    diagrams[6] = render_dot("fig6_disk", diagram_source(
        "MiniFS v2 与内核槽磁盘布局",
        '''
disk [shape=record, style="filled", fillcolor="white",
label="{8 MiB disk.img（16384 × 512 B）|{扇区 0\\n超级块|扇区 1\\ninode 位图|扇区 2-5\\n数据块位图|扇区 6-37\\ninode 表|扇区 38-15359\\nMiniFS 数据区|扇区 15360\\n64 B 内核头|扇区 15361-16383\\nkernel.bin}}"];
note [label="MiniFS：256 inode，单文件最大 64 KiB\\n内核槽：固定加载地址 0x80200000"];
disk -> note [style=dashed, arrowhead=none];
'''))
    diagrams[7] = render_dot("fig7_shell", diagram_source(
        "Shell 执行外部程序的完整链路",
        '''
read [label="Shell 读取并解析命令\\nargv / 重定向 / 后台标志"];
fork [label="fork()", shape=diamond, fillcolor="white"];
parent [label="父进程"];
child [label="子进程"];
redirect [label="open + dup2\\n应用 <、>、>>"];
exec [label="PATH 搜索并 execve"];
elf [label="内核 ELF loader\\n建立用户地址空间与初始栈"];
run [label="sret 进入 U-mode\\ncrt0 → main"];
wait [label="前台：waitpid\\n后台：继续读取命令"];
exit [label="exit → ZOMBIE\\n父进程回收"];
read -> fork;
fork -> parent [label="pid > 0"];
fork -> child [label="pid = 0"];
child -> redirect -> exec -> elf -> run -> exit;
parent -> wait;
exit -> wait;
wait -> read;
'''))
    return diagrams


def capture_terminal():
    command = (
        "cd /home/xiaowen/projects/mycpu && "
        "printf 'help\\nexit\\n' | timeout 20 ./build_wsl/cemu "
        "./os/build/boot.bin --disk ./os/disk.img"
    )
    result = subprocess.run(
        ["wsl", "-d", "Ubuntu", "--", "bash", "-lc", command],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=30,
    )
    output = re.sub(r"\x1b\[[0-9;]*[A-Za-z]", "", result.stdout)
    start = output.find("[BOOT] MiniOS loader")
    if start >= 0:
        output = output[start:]
    lines = output.splitlines()
    selected = []
    for line in lines:
        if len(selected) < 42:
            selected.append(line.rstrip())
        if "Shell exited" in line:
            break
    text = "\n".join(selected)
    if "Welcome to MiniOS" not in text:
        raise RuntimeError("未能捕获 MiniOS 启动输出")

    width, height = 1600, max(720, 56 + len(selected) * 28)
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font_path = Path("C:/Windows/Fonts/consola.ttf")
    if not font_path.exists():
        font_path = Path("C:/Windows/Fonts/lucon.ttf")
    font = ImageFont.truetype(str(font_path), 21)
    draw.rectangle((1, 1, width - 2, height - 2), outline="black", width=3)
    draw.rectangle((1, 1, width - 2, 42), fill="#E6E6E6", outline="black", width=2)
    draw.text((22, 10), "MiniOS / cemu - 实际运行输出", font=font, fill="black")
    y = 56
    for line in selected:
        draw.text((24, y), line, font=font, fill="black")
        y += 28
    path = ASSETS / "fig8_terminal.png"
    image.save(path)
    (ASSETS / "terminal_output.txt").write_text(text, encoding="utf-8")
    return path


def add_figure(doc, path, caption, width=15.6, own_page=False):
    with Image.open(path) as image:
        pixel_width, pixel_height = image.size
    if own_page:
        doc.add_page_break()
    max_height = 20.5 if own_page else 16.5
    width = min(width, max_height * pixel_width / pixel_height) * FIGURE_SCALE
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(path), width=Cm(width))
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.keep_with_next = True
    cp.paragraph_format.space_after = Pt(8)
    set_run_font(cp.add_run(caption), size=10.5)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_CN
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CN)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_LATIN)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    normal.font.size = Pt(12)

    settings = {
        "Heading 1": (FONT_HEADING, 16, 18, 12),
        "Heading 2": (FONT_HEADING, 14, 14, 8),
        "Heading 3": (FONT_HEADING, 12, 10, 6),
    }
    for name, (font, size, before, after) in settings.items():
        style = styles[name]
        style.font.name = font
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_LATIN)
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_LATIN)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = FONT_CN
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CN)
        style.font.size = Pt(12)


def add_front_matter(doc):
    cover = doc.sections[0]
    set_section_geometry(cover)
    cover.top_margin = Cm(2.5)
    v_align = OxmlElement("w:vAlign")
    v_align.set(qn("w:val"), "center")
    cover._sectPr.append(v_align)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(52)
    set_run_font(p.add_run("操作系统课程设计结题报告"), name=FONT_HEADING, size=20, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(70)
    set_run_font(p.add_run(TITLE), name=FONT_HEADING, size=25, bold=True)

    for label, value in (("学生姓名", AUTHOR), ("学号", STUDENT_ID), ("日期", DATE_TEXT)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(14)
        set_run_font(p.add_run(f"{label}：{value}"), size=14)

    front = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_geometry(front)
    front.header.is_linked_to_previous = False
    front.footer.is_linked_to_previous = False
    set_page_numbering(front, "lowerRoman", 1)
    add_page_field(front.footer.paragraphs[0])

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    set_run_font(p.add_run("摘  要"), name=FONT_HEADING, size=16, bold=True)
    abstract = (
        "本项目在自研 RISC-V RV64I 模拟平台 cemu 上设计并实现教学型操作系统 MiniOS。"
        "系统采用两阶段启动方式，由 M 态 Bootloader 从持久化块设备读取、校验并加载内核，"
        "随后切换至 S 态运行。内核实现 Sv39 虚拟内存、物理页分配、Trap 与系统调用、"
        "FCFS/RR 调度、进程创建与回收、同步原语以及持久化 MiniFS v2 文件系统；"
        "用户空间支持 ELF64 程序、argc/argv/envp 初始栈、简化 libc 和交互式 Shell。"
        "Shell 可通过 fork、execve 与 waitpid 执行独立用户程序，并支持目录文件操作、"
        "输入输出重定向、后台任务、ps 与 kill。项目以 98 项 CTest 和 7 项 MiniOS "
        "系统级集成测试验证模拟平台、启动链路、文件系统、进程模型和持久化行为。"
        "结果表明，该系统已经形成从模拟硬件、Bootloader、内核到用户程序的完整运行闭环，"
        "能够用于理解 RISC-V 特权级、虚拟内存和操作系统核心机制之间的协作关系。"
    )
    add_body_paragraph(doc, abstract)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.line_spacing = 1.5
    set_run_font(p.add_run("关键词："), name=FONT_HEADING, size=12, bold=True)
    set_run_font(p.add_run("RISC-V；MiniOS；Sv39；进程调度；MiniFS；ELF"), size=12)

    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    set_run_font(p.add_run("目  录"), name=FONT_HEADING, size=16, bold=True)
    toc = doc.add_paragraph()
    add_toc(toc)

    body = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_geometry(body)
    body.header.is_linked_to_previous = False
    body.footer.is_linked_to_previous = False
    set_page_numbering(body, "decimal", 1)
    header = body.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(header.add_run(SHORT_TITLE), size=9, color=GRAY)
    add_page_field(body.footer.paragraphs[0])


def add_references(doc):
    p = doc.add_paragraph("参考文献", style="Heading 1")
    p.paragraph_format.page_break_before = REFERENCE_PAGE_BREAK
    refs = [
        "[1] RISC-V International. The RISC-V Instruction Set Manual, Volume II: Privileged Architecture.",
        "[2] RISC-V International. The RISC-V Instruction Set Manual, Volume I: Unprivileged ISA.",
        "[3] System V Application Binary Interface. AMD64 Architecture Processor Supplement（ELF 格式参考）.",
        "[4] 《操作系统课程设计》项目制题目（2025-2026 学年春季学期）.",
        "[5] 张晓文. myCPU / MiniOS 项目源代码与开发文档, 2026.",
        "[6] Abraham Silberschatz, Peter B. Galvin, Greg Gagne. Operating System Concepts.",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.74)
        p.paragraph_format.first_line_indent = Cm(-0.74)
        p.paragraph_format.line_spacing = 1.5
        set_run_font(p.add_run(ref), size=10.5)


def parse_body(doc, diagrams, terminal):
    text = SOURCE.read_text(encoding="utf-8")
    first_heading = re.search(r"(?m)^# 1\.", text)
    if first_heading is None:
        raise ValueError("报告源文件中未找到第 1 章标题")
    start = first_heading.start()
    lines = text[start:].splitlines()
    placeholders = {
        "【图 1 占位": (diagrams[1], "图 1  MiniOS 总体分层架构", "【图 1 占位" in OWN_PAGE_FIGURES),
        "【图 2 占位": (diagrams[2], "图 2  cemu 模拟器模块关系", "【图 2 占位" in OWN_PAGE_FIGURES),
        "【图 3 占位": (diagrams[3], "图 3  CPU 单条指令执行流程", "【图 3 占位" in OWN_PAGE_FIGURES),
        "【图 4 占位": (diagrams[4], "图 4  MiniOS 两阶段启动流程", "【图 4 占位" in OWN_PAGE_FIGURES),
        "【图 5 占位": (diagrams[5], "图 5  Trap、系统调用与调度链路", "【图 5 占位" in OWN_PAGE_FIGURES),
        "【图 6 占位": (diagrams[6], "图 6  MiniFS v2 磁盘布局", "【图 6 占位" in OWN_PAGE_FIGURES),
        "【图 7 占位": (diagrams[7], "图 7  Shell 执行外部程序链路", "【图 7 占位" in OWN_PAGE_FIGURES),
        "【截图占位 1": (terminal, "图 8  MiniOS 启动到 Shell 的实际运行输出", False),
    }
    i = 0
    references_added = False
    while i < len(lines):
        line = lines[i].rstrip()
        if any(line.startswith(key) for key in placeholders):
            key = next(key for key in placeholders if line.startswith(key))
            path, caption, own_page = placeholders[key]
            add_figure(
                doc, path, caption,
                width=15.5 if key != "【截图占位 1" else 15.8,
                own_page=own_page,
            )
            i += 1
            continue
        if line.startswith("# 附录") and not references_added and INCLUDE_REFERENCES:
            add_references(doc)
            references_added = True
        if line.startswith("### "):
            doc.add_paragraph(line[4:].strip(), style="Heading 3")
            i += 1
            continue
        if line.startswith("## "):
            doc.add_paragraph(line[3:].strip(), style="Heading 2")
            i += 1
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(line[2:].strip(), style="Heading 1")
            p.paragraph_format.page_break_before = CHAPTER_PAGE_BREAK
            i += 1
            continue
        if line.startswith("```"):
            lang = line[3:].strip()
            block = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            add_code_block(doc, "\n".join(block))
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|?[\s:|-]+\|", lines[i + 1]):
            rows = []
            rows.append([c.strip() for c in line.strip("|").split("|")])
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            add_markdown_table(doc, rows)
            continue
        numbered = re.match(r"^(\d+)\.\s+(.*)$", line)
        if numbered:
            start_value = int(numbered.group(1))
            items = []
            while i < len(lines):
                match = re.match(r"^(\d+)\.\s+(.*)$", lines[i].rstrip())
                if not match:
                    break
                items.append(match.group(2))
                i += 1
            add_numbered_group(doc, items, start_value)
            continue
        if re.match(r"^[-*]\s+", line):
            add_list_item(doc, re.sub(r"^[-*]\s+", "", line), numbered=False)
            i += 1
            continue
        if line.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.74)
            p.paragraph_format.right_indent = Cm(0.74)
            p.paragraph_format.line_spacing = 1.4
            add_inline_markup(p, line.lstrip("> ").strip(), size=10.5, color=GRAY)
            i += 1
            continue
        if line.strip() in ("---", ""):
            i += 1
            continue
        para = [line.strip()]
        i += 1
        while i < len(lines):
            nxt = lines[i].rstrip()
            if (
                not nxt.strip()
                or nxt.startswith(("#", "```", "|", ">", "【"))
                or re.match(r"^\d+\.\s+", nxt)
                or re.match(r"^[-*]\s+", nxt)
                or nxt.strip() == "---"
            ):
                break
            para.append(nxt.strip())
            i += 1
        add_body_paragraph(doc, " ".join(para))
    if not references_added and INCLUDE_REFERENCES:
        add_references(doc)


def set_update_fields(doc):
    settings = doc.settings._element
    node = settings.find(qn("w:updateFields"))
    if node is None:
        node = OxmlElement("w:updateFields")
        settings.append(node)
    node.set(qn("w:val"), "true")


def build():
    ASSETS.mkdir(parents=True, exist_ok=True)
    OUTPUT.mkdir(parents=True, exist_ok=True)
    diagrams = build_diagrams()
    terminal = capture_terminal()

    doc = Document()
    configure_styles(doc)
    add_front_matter(doc)
    parse_body(doc, diagrams, terminal)
    set_update_fields(doc)
    doc.core_properties.title = TITLE
    doc.core_properties.author = AUTHOR
    doc.core_properties.subject = "操作系统课程设计结题报告"
    doc.core_properties.keywords = "RISC-V, MiniOS, Sv39, MiniFS, ELF"
    doc.save(DOCX_OUT)
    print(DOCX_OUT)


if __name__ == "__main__":
    build()
