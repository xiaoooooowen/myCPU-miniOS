from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "操作系统项目进度报告.docx"
ARCH = Path(__file__).resolve().parent / "minios_architecture.png"

INK = "000000"
TEAL = "000000"
BLUE = "000000"
MUTED = "555555"
LIGHT = "F2F2F2"
GOLD = "000000"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size=10.5, bold=False, color=INK, name="Microsoft YaHei"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_text(p, text, **kwargs):
    run = p.add_run(text)
    set_run_font(run, **kwargs)
    return run


def set_para(p, before=0, after=6, line=1.35, align=None, keep=False):
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if keep:
        fmt.keep_with_next = True
    if align is not None:
        p.alignment = align


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    set_para(p, after=7, line=1.42)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix and text.startswith(bold_prefix):
        add_text(p, bold_prefix, bold=True, color=TEAL)
        add_text(p, text[len(bold_prefix):])
    else:
        add_text(p, text)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    set_para(p, after=3, line=1.3)
    add_text(p, text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    set_para(p, after=3, line=1.3)
    add_text(p, text)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    set_para(p, before=13 if level == 1 else 9, after=6, line=1.1, keep=True)
    size = 16 if level == 1 else 12.5
    color = TEAL if level == 1 else BLUE
    add_text(p, text, size=size, bold=True, color=color)
    return p


def add_callout(doc, title, body):
    p = doc.add_paragraph()
    set_para(p, before=2, after=8, line=1.4)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_text(p, title + "：", size=10.5, bold=True)
    add_text(p, body, size=10.5)


def add_progress_table(doc):
    rows = [
        ("第一阶段", "UART/PLIC 输入中断、输入缓冲、sys_read", "终端输入演示与中断日志"),
        ("第二阶段", "用户页表、U 模式切换、用户栈、页异常策略", "两个隔离用户程序"),
        ("第三阶段", "进程状态、fork/exec/wait、锁与信号量", "父子进程和同步演示"),
        ("第四阶段", "RAMFS、文件描述符、ELF 加载、Shell", "ls/cat/echo/ps 综合演示"),
        ("收尾阶段", "回归测试、性能记录、答辩材料和视频", "文档、PPT、演示视频"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(1.05), Inches(3.65), Inches(1.65)]
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, text in enumerate(("阶段", "主要工作", "可检查产出")):
        c = hdr.cells[i]
        c.width = widths[i]
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(c, "E7E7E7")
        set_cell_margins(c)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_para(p, after=0, line=1.15)
        add_text(p, text, size=9.5, bold=True, color=TEAL)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].width = widths[i]
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[i])
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i != 1 else WD_ALIGN_PARAGRAPH.LEFT
            set_para(p, after=0, line=1.2)
            add_text(p, text, size=9.2)
    table.style = "Table Grid"


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    set_run_font(run, size=9, color=MUTED)


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(2.5)
section.bottom_margin = Cm(2.5)
section.left_margin = Cm(2.6)
section.right_margin = Cm(2.6)
section.header_distance = Cm(0.9)
section.footer_distance = Cm(0.8)

sect_pr = section._sectPr
v_align = OxmlElement("w:vAlign")
v_align.set(qn("w:val"), "center")
sect_pr.append(v_align)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(INK)

for style_name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "List Bullet", "List Number"):
    st = styles[style_name]
    st.font.name = "Microsoft YaHei"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

# Cover
p = doc.add_paragraph()
set_para(p, before=0, after=18, line=1.15, align=WD_ALIGN_PARAGRAPH.CENTER)
add_text(p, "基于自研 RISC-V 模拟器的", size=24, bold=True)

p = doc.add_paragraph()
set_para(p, after=72, line=1.15, align=WD_ALIGN_PARAGRAPH.CENTER)
add_text(p, "MiniOS 内核实现", size=32, bold=True)

p = doc.add_paragraph()
set_para(p, after=0, line=1.2, align=WD_ALIGN_PARAGRAPH.CENTER)
add_text(p, "学生姓名：张晓文", size=15)

# Start body in a new, top-aligned section.
body_section = doc.add_section(WD_SECTION.NEW_PAGE)
body_section.page_width = Cm(21)
body_section.page_height = Cm(29.7)
body_section.top_margin = Cm(2.1)
body_section.bottom_margin = Cm(1.9)
body_section.left_margin = Cm(2.6)
body_section.right_margin = Cm(2.6)
body_section.header_distance = Cm(0.9)
body_section.footer_distance = Cm(0.8)
body_v_align = OxmlElement("w:vAlign")
body_v_align.set(qn("w:val"), "top")
body_section._sectPr.append(body_v_align)

body_section.header.is_linked_to_previous = False
header = body_section.header.paragraphs[0]
set_para(header, after=0, line=1)
add_text(header, "操作系统课程设计进度报告", size=8.5, color=MUTED)

body_section.footer.is_linked_to_previous = False
footer = body_section.footer.paragraphs[0]
add_page_number(footer)

# Body content
add_heading(doc, "一、项目概述", 1)
for text in [
    "本项目选择“OS 内核实现”方向。与直接在 QEMU 或 Bochs 上开发内核不同，本项目先使用 C++23 实现一套可控的 RISC-V RV64I 模拟器，再在该模拟器上从零构建 MiniOS。当前系统已经形成从指令执行、总线与设备、特权级切换，到内核启动、中断处理、内存管理和任务调度的完整运行链路。",
    "这种技术路线增加了前期工作量，但也使操作系统与底层硬件之间的关系更加清晰：内核中的一次 printk 最终会转化为 CPU 执行的存储指令，经 MMU 地址翻译和 Bus 路由后写入 UART；一次定时器中断则由 CLINT 产生，经 CSR 与 trap 机制进入内核，最后驱动抢占式调度。项目的核心目标不是堆叠命令或界面，而是建立一个可运行、可追踪、可验证的最小操作系统实验平台。",
]:
    add_body(doc, text)
add_callout(doc, "阶段判断", "截至本报告提交时，项目已稳定完成系统启动、中断与系统调用、内存管理、任务调度四个模块的主要基础功能，共形成 14 个可演示功能点，达到课程“至少 3 个模块、9 个功能点”的基本要求。")

add_heading(doc, "二、总体技术方案", 1)
for text in [
    "项目分为模拟硬件层、MiniOS 内核层和后续用户空间三层。",
    "模拟硬件层由 myCPU 提供，包含 RV64I 指令执行、M/S/U 特权模式、CSR、异常模型、Sv39 MMU、128 MB DRAM、MMIO Bus，以及 UART、CLINT、PLIC 设备模型。cemu 将 kernel.bin 加载到 0x80000000，并从内核入口开始执行。",
    "MiniOS 内核由 C 和 RISC-V 汇编编写。启动代码负责设置栈、清零 BSS、配置异常委托，并通过 mret 从 M 模式进入 S 模式。进入 kernel_main 后，内核依次初始化串口日志、物理页分配器、Sv39 页表、trap 处理、系统调用和任务调度。当前任务在 S 模式运行，已具备协作式与定时器驱动的抢占式切换能力。",
    "后续用户空间将在现有内核之上扩展 U 模式地址空间、ELF 程序加载、进程原语、RAMFS 和 Shell。图中虚线部分表示已经确定接口方向，但尚未计入当前完成功能。",
]:
    add_body(doc, text)

if ARCH.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para(p, before=4, after=3, line=1)
    p.add_run().add_picture(str(ARCH), width=Inches(6.25))
    cap = doc.add_paragraph()
    set_para(cap, after=8, line=1.1, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(cap, "图 1　myCPU + MiniOS 当前架构与下一阶段边界", size=9, color=MUTED)

add_heading(doc, "三、当前阶段成果", 1)
sections = [
    ("3.1 系统启动与运行环境",
     "内核镜像采用独立链接脚本组织，入口地址与模拟器 DRAM 起始地址保持一致。start.S 完成栈指针设置、BSS 清零、medeleg/mideleg 配置、stvec 设置和 M 到 S 特权级切换，随后进入 C 语言内核主函数。启动后 UART 能稳定输出 MiniOS booting... 等日志，多次运行结果一致。",
     ["加载 kernel.bin 到指定物理地址", "设置启动栈并建立 C 运行环境", "初始化 trap 向量", "完成 M 模式到 S 模式切换", "通过 UART 输出启动日志"]),
    ("3.2 中断、异常与系统调用",
     "项目已建立完整的 trap 入口。汇编入口保存 32 个通用寄存器形成 trap frame，C 处理函数读取 scause/sepc/stval 判断异常类型，并在返回前恢复上下文。CLINT 的 mtime/mtimecmp 用于产生周期性定时器中断。系统调用采用 a7 保存调用号、a0-a2 传参、a0 返回结果的约定，当前已实现 sys_write 和 sys_exit 的分派框架。",
     ["trap frame 保存、分派与恢复", "CLINT 周期性定时器中断", "ecall 系统调用入口及 sys_write/sys_exit"]),
    ("3.3 内存管理",
     "物理内存管理采用 bump allocator 与空闲链表组合方案，避免在模拟器中初始化遍历全部物理页。虚拟内存采用 RISC-V Sv39 三级页表，以 2 MB 大页建立 DRAM 身份映射，并映射 UART、CLINT、PLIC 等 MMIO 区域。开启 SATP 后，已有内核功能仍可正常运行。",
     ["4 KB 物理页分配与释放", "释放后页面复用", "Sv39 地址翻译及 DRAM/MMIO 映射"]),
    ("3.4 任务与调度",
     "任务模块实现了 PCB、独立内核栈和上下文结构。协作式调度由 yield 调用 switch_to；抢占式调度利用定时器中断保存的 trap frame，写入下一任务上下文和 sepc，在 sret 后切换执行流。",
     ["PCB 与任务创建", "协作式轮转调度", "定时器驱动的抢占式 Round-Robin 调度"]),
]
for heading, body, points in sections:
    add_heading(doc, heading, 2)
    add_body(doc, body)
    for point in points:
        add_number(doc, point)
    if heading.startswith("3.2"):
        add_callout(doc, "完成边界", "模拟器已能产生并检查页异常，但内核尚未实现按需分配页面，因此“按需分页”不列为当前完成功能。")
    if heading.startswith("3.3"):
        add_body(doc, "当前系统管理 128 MB DRAM，可用页数约 32701 页，明显高于课程要求的 4 MB。")
    if heading.startswith("3.4"):
        add_body(doc, "目前任务状态主要覆盖 UNUSED、READY 和 RUNNING，尚未实现阻塞、僵尸、父子进程关系以及 fork/exec/wait。")

add_heading(doc, "四、测试与阶段验证", 1)
add_body(doc, "模拟器侧使用 Google Test 建立单元测试，当前共 88 项，覆盖 RV64I 指令、CSR、CPU、DRAM、Bus、UART、CLINT、PLIC、异常和 MMU，最近一次冻结记录中全部通过。")
add_body(doc, "内核侧由 cemu 加载 kernel.bin，通过 UART 日志验证以下行为：")
for text in [
    "内核可从固定入口稳定启动并进入 S 模式",
    "物理页分配按 4 KB 对齐，释放页面可再次分配",
    "开启 Sv39 后串口和定时器等 MMIO 访问正常",
    "ecall 可进入 trap handler 并返回",
    "sys_write 能输出字符串，未知系统调用返回错误值",
    "idle、task_a 和 task_b 能在定时器中断驱动下轮转执行",
]:
    add_bullet(doc, text)
add_body(doc, "当前测试主要证明基础机制能够贯通。下一阶段会增加非法地址、资源耗尽、并发竞争等边界场景。")

add_heading(doc, "五、阶段性问题与解决过程", 1)
problems = [
    ("5.1 特权级委托后的中断位映射", "早期 S 模式下定时器中断无法触发调度。原因是 MIP 中的 MTIP 与 SIP 中的 STIP 不在同一 bit，不能简单按位与。项目改为按中断 cause 显式映射挂起位，修复后定时器可以稳定进入 S 模式 trap handler。"),
    ("5.2 trap frame 与 C 函数栈帧混淆", "抢占式调度最初直接在 sched_tick 中读取 sp，但此时 sp 已指向 C 函数自己的栈帧。最终在 trap_entry 中把 trap frame 基址作为参数传给 trap_handler 和 sched_tick，使汇编与 C 之间形成明确的 ABI 契约。"),
    ("5.3 开启分页后外设访问异常", "首次开启 Sv39 时只映射了 DRAM，导致 UART 和 CLINT 的 MMIO 地址触发页异常。修复方案是在内核页表中同时建立 UART、CLINT、PLIC 映射。"),
    ("5.4 裸机环境中的运行库依赖", "内核使用 -nostdlib，64 位除法和取模可能引入 libgcc 辅助函数。printk 的十进制转换改用预计算幂和减法实现，保持纯 RV64I 独立运行。"),
]
for title, body in problems:
    add_heading(doc, title, 2)
    add_body(doc, body)

add_heading(doc, "六、下一阶段四个模块的补全计划", 1)
add_body(doc, "下一阶段不再扩大模拟器基础设施范围，而是围绕“从内核自检走向可交互、可运行用户程序的系统”补全四个模块。")
plans = [
    ("6.1 中断与交互模块",
     "接入 UART 接收中断和 PLIC 外部中断路径，实现字符输入、扫描缓冲和阻塞式读取接口。完成后，系统应能从终端读取一行命令。",
     ["UART 输入经 PLIC 触发外部中断", "输入字符进入内核缓冲区", "提供 read 类系统调用", "连续输入和空缓冲不破坏内核状态"]),
    ("6.2 内存与用户空间模块",
     "为每个进程建立独立页表和用户地址空间，区分代码、数据、堆和栈，并通过 sret 进入 U 模式。页异常处理根据访问类型分配页面或终止进程。",
     ["至少两个进程地址空间相互隔离", "U 模式不能访问内核页", "用户栈可传递参数", "合法缺页可恢复，非法访问只终止当前进程"]),
    ("6.3 进程机制模块",
     "把当前内核任务扩展为进程模型，补充 BLOCKED、ZOMBIE 和父子关系，逐步实现 fork/exec/exit/wait，并以自旋锁和信号量保护共享结构。",
     ["fork 复制基本上下文", "exec 替换用户地址空间", "wait 回收僵尸进程", "阻塞与唤醒不丢失", "RR 时间片可配置"]),
    ("6.4 文件系统与程序加载模块",
     "优先实现 RAMFS，提供 inode、目录项、文件描述符和路径解析；随后实现 ELF 解析器和最小 Shell，使 ls、cat、echo、ps 与外部程序通过统一系统调用运行。",
     ["至少 128 个文件和 3 层目录", "单文件支持读写与 seek", "至少运行 5 个用户程序", "Shell 支持命令与参数解析", "用户程序异常不影响内核"]),
]
for title, body, checks in plans:
    add_heading(doc, title, 2)
    add_body(doc, body)
    p = doc.add_paragraph()
    set_para(p, after=3, line=1.2)
    add_text(p, "验收目标", size=10.2, bold=True, color=GOLD)
    for check in checks:
        add_bullet(doc, check)

add_heading(doc, "七、进度安排", 1)
add_progress_table(doc)
add_body(doc, "执行时以接口贯通为优先：先让输入中断和用户态运行，再引入进程与文件系统。若时间受限，文件系统采用 RAMFS，管道与复杂信号机制作为扩展项，不影响核心验收。")

add_heading(doc, "八、分工方案与 AI 工具使用说明", 1)
add_body(doc, "本项目由一人独立完成，工作范围包括模拟器 C++ 代码、MiniOS 的 C/汇编代码、测试用例、构建脚本和项目文档。")
add_body(doc, "AI 工具主要用于资料检索辅助、代码审查、故障定位思路整理和文档排版。涉及 RISC-V 指令语义、CSR 位定义、页表权限、trap frame 布局等关键内容，均以源码、测试结果和实际运行日志复核。最终代码理解、功能验收和答辩解释由本人负责。")

add_heading(doc, "九、项目创新点", 1)
innovations = [
    "自研 RISC-V 模拟器，使取指、地址翻译、异常进入和设备访问均可逐步追踪。",
    "完整经历 M 模式启动、异常与中断委托、S 模式内核运行的迁移过程。",
    "同时覆盖模拟器 MMU 与内核页表两侧，贯通 Sv39 翻译、权限检查和实际映射。",
    "保留协作式与抢占式两条上下文切换路径，便于对比函数调用上下文与中断上下文。",
]
for index, item in enumerate(innovations, start=1):
    p = doc.add_paragraph()
    set_para(p, after=4, line=1.3)
    add_text(p, f"{index}.　{item}")

add_heading(doc, "十、当前结论", 1)
add_body(doc, "项目已经完成从“能执行 RISC-V 指令”到“能承载 MiniOS 内核”的关键跨越。启动、trap、时钟中断、系统调用、页分配、Sv39 和抢占式调度已经形成可重复运行的闭环。当前重点不再是继续扩充底层模拟器，而是把已有机制组织为真正的用户进程、文件和交互环境。")
add_body(doc, "后续四个模块的完成将使系统从内核功能演示升级为可在自研平台上启动、输入命令、加载程序并管理进程与文件的简化操作系统。")

doc.core_properties.title = "基于自研 RISC-V 模拟器的 MiniOS 内核实现 - 进度报告"
doc.core_properties.subject = "操作系统课程设计第14周进度报告"
doc.core_properties.author = "学生"
doc.core_properties.keywords = "RISC-V, MiniOS, 操作系统, 进度报告"
doc.save(OUT)
print(OUT)
